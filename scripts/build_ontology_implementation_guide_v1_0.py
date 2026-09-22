"""Build the Chinese ontology implementation guide and its diagrams."""
from pathlib import Path
from urllib.parse import unquote
import math
import re
from PIL import Image, ImageDraw, ImageFont
from markdown_it import MarkdownIt
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
NAME = "本体结构与数据解析及智能体驱动说明_v1.0"
MD = ROOT / "docs" / (NAME + ".md")
ASSETS = ROOT / "docs" / "assets" / "本体说明_v1.0"
FONT = "C:/Windows/Fonts/msyh.ttc"
BOLD = "C:/Windows/Fonts/msyhbd.ttc"
BLUE = "#185C97"
INK = "#243746"
MUTED = "#637484"


def font(size, bold=False):
    return ImageFont.truetype(BOLD if bold else FONT, size)


def wrap(draw, text, face, width):
    lines = []
    for part in text.split("\n"):
        line = ""
        for char in part:
            if draw.textlength(line + char, font=face) > width and line:
                lines.append(line)
                line = char
            else:
                line += char
        lines.append(line)
    return lines


def label(draw, rect, title, sub="", fill="#F0F6FC", border="#BDD1E5"):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=10, fill=fill, outline=border, width=2)
    lines = [(s, font(28, True), INK) for s in wrap(draw, title, font(28, True), x2-x1-32)]
    if sub:
        lines += [(s, font(22), MUTED) for s in wrap(draw, sub, font(22), x2-x1-32)]
    heights = [f.size+12 for _, f, _ in lines]
    y = y1 + (y2-y1-sum(heights))/2
    for (s, f, color), height in zip(lines, heights):
        draw.text(((x1+x2-draw.textlength(s, font=f))/2, y), s, font=f, fill=color)
        y += height


def arrow(draw, start, end, color=BLUE):
    draw.line([start, end], fill=color, width=4)
    theta = math.atan2(end[1]-start[1], end[0]-start[0])
    points = [end]
    for offset in (-0.55, 0.55):
        points.append((end[0]-17*math.cos(theta+offset), end[1]-17*math.sin(theta+offset)))
    draw.polygon(points, fill=color)


def canvas(title, subtitle, height):
    image = Image.new("RGB", (1600, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 10), fill=BLUE)
    draw.text((48, 35), title, font=font(36, True), fill=INK)
    draw.text((48, 94), subtitle, font=font(22), fill=MUTED)
    return image, draw


def diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    image, draw = canvas("01  本体、知识与业务的分层关系", "现有存储与工具并存；语义声明不等于业务已经自动同步", 1030)
    rows = [
        ("智能体应用", "数据处理智能体  /  对话助手  /  六个营销智能域", "共用 Harness；各入口实际读取、执行能力不同"),
        ("业务与图谱实例", "活动、产品包、审批等业务记录  |  本体实体与关系", "两类存储并存；完整业务变更投影仍需补齐"),
        ("业务语义注册", "41 类对象  /  44 类关系  /  14 个动作  /  13 个函数", "定义对象和关联；动作、函数主要为语义声明"),
        ("知识与证据", "知识文档  /  文本片段  /  来源  /  抽取任务与事件", "保存原文依据；精确页码、证据定位仍有限"),
        ("数据输入", "上传文档  /  接口快照  /  航班产品适配  /  市场热点", "不同入口有不同处理路径；不等同于通用数据库同步"),
    ]
    for i, (title, body, note) in enumerate(rows):
        y = 166+i*162
        label(draw, (48,y,318,y+124), title, fill="#E4EFFA")
        label(draw, (340,y,1552,y+124), body, note, fill="#F8FAFC")
        if i < len(rows)-1:
            arrow(draw, (945,y+160), (945,y+129))
    image.save(ASSETS / "01_结构.png")

    image, draw = canvas("02  文件到知识与本体的处理链路", "主要文件入口的实际流程；接口与航班产品适配存在不同实现", 1120)
    xs = [(48,488),(580,1020),(1112,1552)]
    for (x1,x2), title, sub in zip(xs, ["上传文件","解析为文本","保存知识"], ["格式与大小检查 / 创建任务","MinerU 或本地文本解析","文档、版本、1200 字符片段"]):
        label(draw, (x1,178,x2,320), title, sub)
    arrow(draw,(490,249),(574,249)); arrow(draw,(1022,249),(1106,249))
    arrow(draw,(1332,325),(1332,391))
    for (x1,x2), title, sub in zip(xs, ["本体准入校验","候选对象与关系","数据处理智能体"], ["类型 / 置信度 / 证据 / 关系端点","抽取决策、属性、引用与理由","文件名、正文前 5 万字符、类型约定"]):
        label(draw,(x1,399,x2,551),title,sub)
    arrow(draw,(1106,475),(1027,475)); arrow(draw,(574,475),(494,475))
    arrow(draw,(268,556),(268,644))
    label(draw,(48,651,488,830),"不适合更新本体","仅保存知识，保留抽取结果",fill="#F1F5F8")
    draw.line([(400,556),(400,603),(800,603)],fill=BLUE,width=4)
    arrow(draw,(800,603),(800,644))
    draw.text((548,566),"适合更新",font=font(22),fill=BLUE)
    label(draw,(580,651,1020,830),"人工二次确认","批准后写入；驳回仅保留知识",fill="#FFF8E8",border="#DCC580")
    arrow(draw,(1024,740),(1107,740))
    label(draw,(1112,651,1552,830),"写入本体","实体更新 / 关系写入 / 证据关联",fill="#EDF8F2",border="#9ECDB3")
    label(draw,(48,910,1552,1064),"处理事件与来源贯穿任务","当前限制：实体准入过滤存在缺口；证据多关联首片段；非逐候选审批。详见正文第 5、7、14 章。",fill="#FFF5EF",border="#E1BEA7")
    image.save(ASSETS / "02_流水线.png")

    image, draw = canvas("03  智能体如何读取并使用业务知识", "当前工具调用链与语义契约执行边界", 1150)
    label(draw,(390,164,1210,274),"业务人员提出问题或提交任务","问题、活动标识、租户上下文")
    arrow(draw,(800,280),(800,333))
    label(draw,(270,340,1330,484),"统一 Harness + 对话规划器","模型调用、工具执行、运行事件、用量记录；规划最多 4 轮")
    for x in (275,800,1325):
        arrow(draw,(x,488),(x,561))
    label(draw,(48,568,502,752),"检索工具","知识检索 / 本体检索",fill="#EEF6FD")
    label(draw,(574,568,1026,752),"业务查询工具","活动检查 / 产品查询 / 任务检查",fill="#EDF8F2")
    label(draw,(1098,568,1552,752),"六智能域运行器","机会 / 客群 / 匹配 / 编排 / 内容 / 效果",fill="#FFF8E8")
    label(draw,(48,809,502,984),"知识与图谱存储","当前以关键词、实体标签查询为主")
    label(draw,(574,809,1026,984),"业务记录与处理结果","按租户读取；产品查询不等于实时库存")
    label(draw,(1098,809,1552,984),"契约已定义，执行待完善","尚未自动装载子图与执行完整业务函数")
    for x in (275,800,1325):
        arrow(draw,(x,758),(x,803))
    draw.text((48,1040),"工具结果返回模型：汇总建议 + 来源；业务写入仍须遵循各入口审批与权限。",font=font(26,True),fill=INK)
    image.save(ASSETS / "03_智能体.png")


def eastasia(style, size, bold=False, color=None):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color):
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(item)


def inline(paragraph, token):
    strong = False
    emphasis = False
    for item in token.children or []:
        if item.type == "strong_open": strong = True
        elif item.type == "strong_close": strong = False
        elif item.type == "em_open": emphasis = True
        elif item.type == "em_close": emphasis = False
        elif item.type in ("softbreak", "hardbreak"): paragraph.add_run().add_break()
        elif item.type in ("text", "code_inline"):
            run = paragraph.add_run(item.content)
            run.bold = strong
            run.italic = emphasis
            if item.type == "code_inline":
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("185C97")


def field(paragraph, name):
    run = paragraph.add_run()
    node = OxmlElement("w:fldSimple")
    node.set(qn("w:instr"), name)
    run._r.addnext(node)


def build_word():
    source = MD.read_text(encoding="utf-8").replace("必須", "必须")
    tokens = MarkdownIt("commonmark").enable("table").parse(source)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
    for name,size,bold,color in [("Normal",10.5,False,"243746"),("Title",24,True,"185C97"),("Subtitle",13,False,"637484"),("Heading 1",17,True,"185C97"),("Heading 2",13,True,"185C97"),("Heading 3",11,True,"243746"),("Caption",9,False,"637484")]:
        eastasia(doc.styles[name],size,bold,color)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    doc.styles["Heading 1"].paragraph_format.page_break_before = False
    for name in ("Heading 1","Heading 2","Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
    eastasia(doc.styles["Header"],8,False,"637484")
    sec.header.paragraphs[0].text = "东方航空智能营销平台  |  本体与智能体技术说明  v1.0"
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("第 "); field(footer,"PAGE"); footer.add_run(" 页")
    doc.add_paragraph("东方航空智能营销平台", "Title")
    doc.add_paragraph("本体结构与数据解析\n及智能体驱动说明", "Title")
    doc.add_paragraph("v1.0  |  2026年9月20日", "Subtitle")
    doc.add_paragraph("基于当前实现的业务与技术说明")
    doc.add_paragraph("覆盖本体语义、数据解析、知识与证据、人工确认、营销业务关联、六智能域契约、工具调用及生产完善建议。")
    doc.add_paragraph("阅读提示：代码中已经声明的对象、动作和函数，不等于已经形成完整的运行闭环。本文对现有能力与待完善项分别说明。")
    doc.add_page_break()
    doc.add_heading("阅读导航",level=1)
    for i,tok in enumerate(tokens):
        if tok.type == "heading_open" and tok.tag == "h2":
            doc.add_paragraph(tokens[i+1].content)
    doc.add_page_break()
    i=0
    list_depth=0
    quote_depth=0
    while i < len(tokens):
        tok=tokens[i]
        if tok.type == "heading_open":
            level=int(tok.tag[1:])
            if level != 1:
                p=doc.add_heading(level=min(level-1,3)); inline(p,tokens[i+1])
                if level == 2 and tokens[i+1].content.startswith(("2. ","5. ","9. ","14. ","附录")):
                    p.paragraph_format.page_break_before=True
            i+=3; continue
        if tok.type == "table_open":
            rows=[]; row=[]; i+=1
            while tokens[i].type != "table_close":
                t=tokens[i]
                if t.type == "tr_open": row=[]
                elif t.type == "inline": row.append(t)
                elif t.type == "tr_close": rows.append(row)
                i+=1
            table=doc.add_table(rows=0,cols=max(map(len,rows)))
            table.style="Table Grid"
            table.autofit=False
            for col in table.columns: col.width=Cm(17.4/len(table.columns))
            for ri,row in enumerate(rows):
                cells=table.add_row().cells
                pr=table.rows[-1]._tr.get_or_add_trPr()
                pr.append(OxmlElement("w:cantSplit"))
                if ri==0: pr.append(OxmlElement("w:tblHeader"))
                for cell,content in zip(cells,row):
                    cell.vertical_alignment=1
                    p=cell.paragraphs[0]; inline(p,content)
                    p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.space_before=Pt(4)
                    p.paragraph_format.line_spacing=1.05
                    for run in p.runs:
                        run.font.size=Pt(9)
                        if ri==0: run.bold=True
                    if ri==0: shade(cell,"E5EFF9")
                    elif ri%2==0: shade(cell,"F5F8FB")
            doc.add_paragraph()
        elif tok.type == "paragraph_open":
            child=tokens[i+1]
            pictures=[x for x in child.children or [] if x.type=="image"]
            if pictures:
                for pic in pictures:
                    path=MD.parent / unquote(pic.attrGet("src"))
                    doc.add_picture(str(path),width=Cm(17.3))
                    doc.paragraphs[-1].alignment=1
                    doc.add_paragraph(pic.content,"Caption")
            else:
                p=doc.add_paragraph(style="List Bullet" if list_depth else "Normal")
                inline(p,child)
                if quote_depth:
                    p.paragraph_format.left_indent=Cm(.4)
            i+=3; continue
        elif tok.type in ("bullet_list_open","ordered_list_open"): list_depth+=1
        elif tok.type in ("bullet_list_close","ordered_list_close"): list_depth-=1
        elif tok.type == "blockquote_open": quote_depth+=1
        elif tok.type == "blockquote_close": quote_depth-=1
        elif tok.type in ("fence","code_block"):
            code_lines=tok.content.rstrip().splitlines()
            for line_index,line in enumerate(code_lines):
                p=doc.add_paragraph()
                p.paragraph_format.keep_with_next=line_index < len(code_lines)-1
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1
                run=p.add_run(line)
                run.font.name="Consolas"; run.font.size=Pt(8)
        i+=1
    doc.core_properties.title="东方航空智能营销平台：本体结构与数据解析及智能体驱动说明 v1.0"
    doc.core_properties.subject="当前代码实现、运行边界与生产完善建议"
    doc.core_properties.author="东方航空智能营销平台项目组"
    out=MD.with_suffix(".docx")
    doc.save(out)
    print(out)
    print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; diagrams: {len(doc.inline_shapes)}")


if __name__ == "__main__":
    diagrams()
    build_word()
