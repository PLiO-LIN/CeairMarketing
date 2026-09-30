"""Create V1.8 with a concise innovation summary and a detailed appendix.

V1.8 intentionally treats section 2.2 as the evaluator's quick-reading layer.
The mechanisms, evidence, comparison points, and implementation boundaries
remain in Appendix A rather than being repeated in the main text.
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "output" / "丽娃谷风_东方航空智能营销平台_项目文档_V1.7_70页.docx"
OUT = ROOT / "output" / "丽娃谷风_东方航空智能营销平台_项目文档_V1.8_70页.docx"


def set_run_font(run, size=11, bold=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def insert_after(anchor, text, style, body=False):
    """Create a styled paragraph immediately after ``anchor``."""
    document = anchor._parent
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)
    if body:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.5
        for run in paragraph.runs:
            set_run_font(run, size=11)
    else:
        paragraph.paragraph_format.space_before = Pt(7)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.25
        for run in paragraph.runs:
            set_run_font(run, size=12, bold=True)
    anchor._p.addnext(paragraph._p)
    return paragraph


def replace_innovation_summary(document):
    paragraphs = document.paragraphs
    heading = "2.2  技术创新点"
    next_section = "3  实施方案"
    start = [i for i, p in enumerate(paragraphs) if p.text.strip() == heading][-1]
    end = next(i for i, p in enumerate(paragraphs[start + 1:], start + 1)
               if p.text.strip() == next_section)

    anchor = paragraphs[start]
    for paragraph in paragraphs[start + 1:end]:
        paragraph._element.getparent().remove(paragraph._element)

    anchor = insert_after(
        anchor,
        "本节仅给出评审可快速核对的创新结论和业务价值；问题拆解、技术机制、工程依据、对比与适用边界统一收录于附录 A，避免主文重复展开。",
        "Normal",
        body=True,
    )

    items = [
        (
            "2.2.1  本体化营销决策上下文",
            "项目把机会、客群快照、产品包、活动、反馈与复盘等对象及其关系组织为营销本体，并将来源证据、有效期、置信度、冲突和确认状态纳入同一语义层。模型据此读取的是可定位、可校验的业务事实，而不是仅凭相似文本拼接建议；过期、冲突或待审信息不能直接支撑可执行动作。详见附录 A.1-A.2。",
        ),
        (
            "2.2.2  统一 Harness 的受控推理运行",
            "统一 Harness 在一次 AI 运行中编排身份与租户校验、授权上下文加载、模型和白名单工具调用、结构与规则校验、事件记录、重试和降级。它将模型调用从页面中的黑箱请求转化为可观察、可回放的业务运行对象，并明确模型生成、确定性规则和人工门禁之间的责任边界。详见附录 A.3。",
        ),
        (
            "2.2.3  持久化人机协同与任务级模型路由",
            "项目将审批、采纳、修改和驳回建模为带版本、角色、理由与状态变化的持久化决策数据；同时以任务而非模型名称定义双模型对照，在相同权限、上下文、工具、输出 Schema 与审核规则下评价质量、时延、成本和安全。达到阈值的任务可逐步扩展，异常时保留受控回退或人工接管。详见附录 A.4-A.5。",
        ),
        (
            "2.2.4  状态可控、可回放的营销活动闭环",
            "平台把机会、客群、产品、内容、活动版本、审批、执行、反馈与复盘串成受状态机约束的对象链；审批冻结版本，反馈和复盘关联批次与指标口径。数据缺失时只保留待补充结论，不将演练数据作经营归因。详见附录 A.6。",
        ),
    ]
    for title, text in items:
        anchor = insert_after(anchor, title, "Heading 3")
        anchor = insert_after(anchor, text, "Normal", body=True)


def update_release_metadata(document):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "V1.7  实质扩展版":
            paragraph.clear()
            run = paragraph.add_run("V1.8  摘要-附录结构优化版")
            set_run_font(run, size=14, bold=True)
        elif paragraph.text.startswith("说明：V1.7 保留赛事模板"):
            paragraph.clear()
            run = paragraph.add_run(
                "说明：V1.8 保留赛事模板规定的项目概况、项目规划、实施方案与参考资料结构。"
                "本次调整将正文 2.2 收束为评审摘要，并以明确索引连接附录 A 的创新论证，"
                "使主文阅读节奏与证据材料深度各自清晰。"
            )
            set_run_font(run, size=10.5)

    history = next(
        table for table in document.tables
        if table.rows and table.cell(0, 0).text.strip() == "序号"
    )
    source_row = history.rows[-1]
    row = history.add_row()
    values = [
        "6", "重构技术创新点的主文与附录层次", "V1.8", "丽娃谷风",
        "2026.09.01", "正文摘要，附录保留完整论证",
    ]
    for index, value in enumerate(values):
        target = row.cells[index]
        source = source_row.cells[index]
        if source._tc.tcPr is not None:
            target._tc.remove(target._tc.tcPr)
            target._tc.insert(0, deepcopy(source._tc.tcPr))
        target.text = value
        for paragraph in target.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, size=10.2)

    document.core_properties.title = "东方航空智能营销平台项目文档 V1.8"
    document.core_properties.comments = "V1.8：正文创新摘要与附录深度论证分层"


def build():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Missing V1.7 source document: {SOURCE}")
    document = Document(SOURCE)
    replace_innovation_summary(document)
    update_release_metadata(document)
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
