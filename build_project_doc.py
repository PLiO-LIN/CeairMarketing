from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "丽娃谷风_东方航空智能营销平台_项目文档_V1.3.docx"

INK = "000000"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
RULE = "A6A6A6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_dxa):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, name="宋体", size=11, bold=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, before=0, after=8, line=1.333, align=None, first_indent=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_indent:
        fmt.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align


def add_text(doc, text, size=11, bold=False, before=0, after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent=True, font="宋体", color=INK):
    p = doc.add_paragraph()
    set_paragraph_format(p, before, after, 1.333, align, first_indent)
    set_run_font(p.add_run(text), font, size, bold, color)
    return p


def add_heading(doc, text, level):
    style_name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.keep_with_next = True
    set_paragraph_format(p, before=18 if level == 1 else 10 if level == 2 else 6, after=8 if level == 1 else 6 if level == 2 else 4,
                         line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
    set_run_font(p.add_run(text), "黑体", 16 if level == 1 else 13 if level == 2 else 11, True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=3, after=8, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    set_run_font(p.add_run(text), "宋体", 9.5, False, GRAY)
    return p


def populate_table(table, rows, header=True):
    for r_idx, values in enumerate(rows):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            set_paragraph_format(p, before=0, after=0, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
            set_run_font(p.add_run(value), "宋体", 9.5, bool(header and r_idx == 0))
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)


def add_table(doc, rows, widths_dxa, header=True):
    table = doc.add_table(rows=len(rows), cols=len(widths_dxa))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    populate_table(table, rows, header=header)
    if header:
        set_repeat_table_header(table.rows[0])
    return table


def add_references_table(doc, references):
    """Keep the short reference list compact without reducing its readability."""
    rows = []
    midpoint = (len(references) + 1) // 2
    for left, right in zip(references[:midpoint], references[midpoint:]):
        rows.append([left, right])
    if len(references) % 2:
        rows.append([references[-1], ""])
    table = add_table(doc, rows, [4680, 4680], header=False)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_paragraph_format(paragraph, before=0, after=1, line=1.12,
                                     align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 9.5)
    return table


def add_page_number(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.color.rgb = RGBColor(0, 0, 0)
    if "Caption" not in doc.styles:
        caption = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "宋体"
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def configure_section(section, first_page=False):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = first_page
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_format(header_p, before=0, after=0, line=1.0, first_indent=False)
    set_run_font(header_p.add_run("东方航空智能营销平台  项目文档"), "宋体", 9, False, GRAY)
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(footer_p, before=0, after=0, line=1.0, first_indent=False)
    set_run_font(footer_p.add_run("第 "), "宋体", 9, False, GRAY)
    add_page_number(footer_p)
    set_run_font(footer_p.add_run(" 页"), "宋体", 9, False, GRAY)


def cover(doc):
    for _ in range(7):
        add_text(doc, "", before=0, after=0, first_indent=False)
    add_text(doc, "第八届中国研究生人工智能创新大赛", size=18, bold=True,
             before=0, after=22, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, font="黑体")
    add_text(doc, "东方航空智能营销平台", size=28, bold=True,
             before=0, after=16, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, font="黑体")
    add_text(doc, "项目文档", size=20, bold=True,
             before=0, after=44, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, font="黑体")
    metadata = [
        ["版本号", "V1.3"],
        ["完成日期", "2026.08.31"],
        ["团队名称", "丽娃谷风"],
        ["团队成员", "胡于飞、林李洋、张婷婷"],
    ]
    table = add_table(doc, metadata, [2600, 6760], header=False)
    for row in table.rows:
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[0].paragraphs[0].runs[0].bold = True
    doc.add_page_break()


def toc(doc):
    add_text(doc, "目录", size=20, bold=True, before=12, after=22,
             align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, font="黑体")
    entries = [
        "1  项目概述",
        "  1.1  背景与意义",
        "    1.1.1  业务痛点",
        "  1.2  项目价值",
        "    1.2.1  服务对象与使用方式",
        "  1.3  资源支持",
        "2  项目规划",
        "  2.1  建设目标",
        "    2.1.1  分阶段建设目标",
        "  2.2  主要创新点",
        "  2.3  双模型演进策略",
        "3  实施方案",
        "  3.1  环境与可行性分析",
        "  3.2  实施细节",
        "  3.3  AI 推理与营销本体协同",
        "  3.4  计划与分工",
        "  3.5  验证指标与典型场景",
        "  3.6  平台功能与业务工作流",
        "  3.7  数据处理与知识本体构建",
        "  3.8  推理编排、模型路由与可观测性",
        "  3.9  AI 评测体系与稳健性验证",
        "  3.10  治理边界与后续建设路线",
        "    3.10.1  典型场景 AI 推理执行清单",
        "4  参考资料",
    ]
    for entry in entries:
        add_text(doc, entry, size=11, before=0, after=5, align=WD_ALIGN_PARAGRAPH.LEFT,
                 first_indent=False)
    add_text(doc, "注：目录页码可在 Microsoft Word 中使用“更新域”同步更新。", size=9.5,
             before=10, after=0, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, color=GRAY)
    doc.add_page_break()


def revision_history(doc):
    add_text(doc, "修订历史", size=18, bold=True, before=12, after=18,
             align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, font="黑体")
    rows = [
        ["编号", "修订原因", "版本", "日期", "修订人", "说明"],
        ["1", "初版编制", "V1.0", "2026.08.30", "丽娃谷风", "根据当前项目实现、技术文档与竞赛模板编制"],
        ["2", "内容扩展与模型说明", "V1.1", "2026.08.31", "丽娃谷风", "补充本地 Qwen 自研模型、双模型协同及实施细节"],
        ["3", "模型路线说明修订", "V1.2", "2026.08.31", "丽娃谷风", "明确前期以百炼为主，后续逐步扩大内置模型使用"],
        ["4", "功能与 AI 技术细化", "V1.3", "2026.08.31", "丽娃谷风", "补充平台功能、数据管道、推理编排、评测体系与治理路线"],
    ]
    add_table(doc, rows, [700, 1900, 900, 1400, 1500, 2960])
    add_text(doc, "文档说明：本文以当前原型实现、团队自研模型能力及项目技术资料为依据。页面展示的营销活动数值属于原型演示数据，不作为生产经营结果或参赛效果承诺。", size=10,
             before=18, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True, color=GRAY)
    doc.add_page_break()


def main_body(doc):
    add_heading(doc, "1  项目概述", 1)
    add_heading(doc, "1.1  背景与意义", 2)
    add_text(doc, "航空营销正从单点促销走向覆盖市场洞察、客群识别、产品组合、内容编排、渠道执行与经营复盘的全链路运营。实际工作中，航线经营、航班运行、用户画像、产品规则、市场热点和渠道回执往往分散在不同系统中，业务人员需要在多源信息之间反复核验，难以快速形成既可执行又可追溯的营销方案。与此同时，生成式人工智能在文本生成和信息归纳方面具备优势，但若缺少业务知识、事实证据和人工审核，容易产生不符合库存、资格、频控或合规约束的建议。")
    add_text(doc, "“东方航空智能营销平台”面向航空公司营销、产品、运营、审批和分析人员，构建以业务本体和可追溯证据为底座的智能营销工作台。平台将市场与经营信号、聚合客户画像、航班及产品信息、内容资产、渠道反馈和复盘结果连接为可治理闭环，让智能体在受控业务语境内提供机会分析、客群洞察、产品匹配、活动编排、内容生成与效果分析建议。")
    add_heading(doc, "1.1.1  业务痛点", 3)
    add_text(doc, "第一，营销线索分散。热点、航线供给、预售、价格、客座率、客户旅程、产品库存和渠道反馈在不同业务系统中产生，数据粒度、更新频率与口径并不一致。传统依赖人工检索和跨表核对的方式，很难在机会窗口内完成可靠的判断。", before=0, after=6)
    add_text(doc, "第二，活动决策链条长。一项航线营销活动需要同时回答“为什么做、面向谁做、卖什么、何时触达、通过什么渠道、如何控制风险、如何评估效果”等问题。任一环节缺乏事实依据，都会导致内容失真、触达不当或资源浪费。", before=0, after=6)
    add_text(doc, "第三，AIGC 的可用性需要业务约束。通用模型可以快速总结和生成文本，但不能天然理解航班、库存、运价、会员权益、旅客授权、频控与审批规则。项目将 AI 推理嵌入事实检索、工具调用、结构化校验和人工门禁之中，使模型从“能生成”走向“能解释、可审核、可复盘”。", before=0, after=8)
    add_heading(doc, "1.2  项目价值", 2)
    value_rows = [
        ["价值维度", "具体体现"],
        ["业务价值", "缩短从市场信号到活动方案的分析链路，帮助运营人员在客群、产品、渠道与时机之间建立可解释关联。"],
        ["管理价值", "将预算、产品可售、客户授权、触达频控、内容事实和敏感词等检查前置，并保留审批、版本、证据和操作审计。"],
        ["技术价值", "以营销业务本体统一表达机会、客群快照、产品包、活动、触点、反馈和归因，支持跨数据源的语义查询与可追溯生成。"],
        ["数据安全价值", "只将聚合客群、画像快照和统计指标用于营销本体，不将个人旅客明细直接写入知识底座；通过租户隔离和权限控制降低数据越权风险。"],
    ]
    add_table(doc, value_rows, [1900, 7460])
    add_heading(doc, "1.2.1  服务对象与使用方式", 3)
    add_text(doc, "营销运营人员可在机会工作台查看 AI 汇总的市场、航线和客户信号，并通过自然语言圈选、产品匹配、内容生成和活动编排缩短方案形成时间；产品与内容人员可针对可售资格、权益规则和渠道规范进行人工审核；审批人员在统一界面核验预算、频控、保护排除和合规条件；分析人员则基于触达、点击、出票、履约和收入回流开展归因与策略学习。", before=0, after=6)
    add_text(doc, "平台强调“人机协同”而非“自动发布”。模型可以负责信息归纳、候选方案生成、风险提示和复盘建议；涉及正式活动发布、本体写入、预算审批、渠道触达和规则生效的操作，仍由具备权限的业务人员确认并留痕。", before=0, after=8)
    add_heading(doc, "1.3  资源支持", 2)
    add_text(doc, "项目已形成可本地运行的原型系统：前端为 `apps/web-v32` 生产工作台页面，后端采用 FastAPI，业务数据与本体关系存储于 PostgreSQL，整体通过 Docker Compose 编排。当前管理工作台已接入阿里云百炼 AgentStudio 的 OpenAI 兼容模型服务（qwen3.7-flash），并由统一 Harness 负责工具调用、结构化决策、来源记录与错误治理。")
    add_text(doc, "除云端模型服务外，团队还具备基于 Qwen 的本地自研模型能力，形成“云端百炼主用、内置模型渐进扩展”的双模型技术路线。当前阶段以已接入的百炼模型服务提供通用推理能力；本地模型先用于受控验证、领域知识理解、结构化抽取和业务辅助等适配场景，后续再依据实测效果逐步扩大使用范围。两类模型均应纳入统一的身份认证、权限范围、日志审计、提示模板、输出 Schema 和人工审核约束，避免模型能力与业务治理脱节。")
    add_text(doc, "当前原型已具备文件接入、知识存储、本体存储、模型配置、营销智能助手、客户画像目录、航班及产品数据标准化处理等基础能力。数据接入和生产发布仍遵循授权、脱敏、人工确认与业务审批的边界，避免把原型演示能力等同于已上线生产能力。")

    add_heading(doc, "2  项目规划", 1)
    add_heading(doc, "2.1  建设目标", 2)
    add_text(doc, "本项目的总体目标是完成一个可演示、可运行、可审计的航空智能营销原型，验证“数据与知识沉淀 - 智能分析 - 人工决策 - 营销执行 - 效果回流”的闭环可行性。系统以营销活动全生命周期为主线，帮助业务人员完成从机会发现到复盘优化的协同工作，而不是以大模型直接替代业务审批或自动触达客户。")
    goal_rows = [
        ["目标层级", "目标内容", "验证方式"],
        ["业务闭环", "贯通机会、客群、产品、内容、审批、执行与复盘七个环节。", "通过工作台页面、数据对象和流程演示验证。"],
        ["智能能力", "形成六个 AI 智能域，输出有来源、有约束、可解释的分析和建议。", "通过统一 Harness 轨迹、知识/本体证据和人工门禁验证。"],
        ["工程能力", "完成前后端、数据库与容器化部署，并对模型调用异常和边界输入进行保护。", "通过自动化测试、健康检查和端到端低流量调用验证。"],
        ["治理能力", "落实租户隔离、角色权限、数据来源、版本留痕与人工确认。", "通过数据模型、接口约束和审计记录验证。"],
    ]
    add_table(doc, goal_rows, [1500, 4350, 3510])
    add_heading(doc, "2.1.1  分阶段建设目标", 3)
    add_text(doc, "第一阶段完成业务对象统一表达。围绕市场信号、机会、客户需求、客群快照、产品包、内容、活动、执行批次、反馈、归因和复盘定义可关联的数据对象，并为每个对象保留来源、时间、置信度、版本与租户标识。", before=0, after=6)
    add_text(doc, "第二阶段完成 AI 推理闭环。模型在接收营销问题后，先加载当前租户允许访问的知识、本体、活动与授权工具上下文，再规划下一步业务工具，完成机会检索、客群解释、产品资格核验、内容生成或效果分析，并将最终答案与证据和待确认事项一并返回。", before=0, after=6)
    add_text(doc, "第三阶段完成可观测与可评测能力。对每次推理记录模型服务、工具调用、上下文范围、结构化输出校验、错误恢复、耗时和人工反馈，使项目既能展示 AI 生成效果，也能分析其业务有效性、可靠性和治理边界。", before=0, after=8)
    add_heading(doc, "2.2  主要创新点", 2)
    innovations = [
        ("（1）面向航空营销闭环的六智能域协同", "平台不是通用问答机器人，而是将机会洞察、客群洞察、产品匹配、活动编排、内容生成和效果分析映射到营销业务链条。智能输出定位为候选方案与解释，关键动作必须经过人工确认。"),
        ("（2）证据驱动的营销本体与知识底座", "以市场信号、航线、航班、客户需求、客群快照、产品包、活动、反馈和复盘等对象及关系表达营销语义，并为对象保留来源、证据、置信度、有效期、版本和租户信息。"),
        ("（3）多源数据的分级准入机制", "文件、接口、数据库和热点信息先经过解析、清洗、实体对齐、关系校验和置信度判断；可稳定映射的事实进入本体候选，低置信度信息进入人工确认或仅保留在知识层。"),
        ("（4）业务约束内置的活动编排", "活动触达前联动产品可售与权益资格、客户授权与保护名单、频控、预算、内容事实与敏感词检查，避免将模型生成内容直接转化为真实营销动作。"),
        ("（5）可诊断的模型调用治理", "统一 Harness 记录工具调用、来源和决策过程，并对网络抖动、响应结构异常、非法输入、租户越权和提示注入等场景实施重试、拒绝、脱敏与恢复策略。"),
        ("（6）本地与云端双模型协同", "团队自研本地 Qwen 模型承载可本地部署的领域推理与结构化处理能力，已接入的百炼模型服务承载云端推理能力。项目通过统一的模型接入和治理框架，为不同业务场景选择合适推理资源，并保持相同的权限、审计和输出约束。"),
        ("（7）以业务结果反向校验 AI 推理", "项目不以“生成了一段文本”作为完成标准，而是将机会置信度、产品适配率、内容点击、转化漏斗、渠道回执、人工采纳与异常补偿纳入复盘链路，形成可量化的推理效果证据。"),
    ]
    for title, detail in innovations:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=3, after=4, line=1.333, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True)
        set_run_font(p.add_run(title), "宋体", 11, True)
        set_run_font(p.add_run(detail), "宋体", 11)

    add_heading(doc, "2.3  双模型演进策略", 2)
    add_text(doc, "项目采用“前期云端主用、后期内置模型逐步扩大”的双模型演进路线。前期暂以阿里云百炼 AgentStudio 的 qwen3.7-flash 作为默认模型服务，快速验证营销智能体、知识检索、本体查询、工具编排和人工审核门禁的端到端链路。该阶段重点不是追求单一模型替换，而是先将业务流程、数据边界、输出格式和评价口径固化下来。")
    add_text(doc, "团队自研的本地 Qwen 模型在当前本地原型中以“内置测试模型（ceair-governed-mock-v1）”登记并启用，用于受控推理、流程对照和治理验证。模型服务层对每次调用统一记录模型名称、提示词令牌数、生成令牌数、总令牌数、工具轨迹和异常状态；因此，后续接入真实本地推理服务时不需要改变上层业务流程，只需在统一提供者接口下配置服务地址、模型标识与权限范围。")
    model_rows = [
        ["阶段", "模型使用策略", "主要目标", "治理要求"],
        ["阶段一：云端主用", "暂以百炼 qwen3.7-flash 作为默认推理服务；内置模型用于受控测试与对照。", "验证六智能域、工具编排、知识/本体上下文和人工门禁。", "模型服务白名单、租户隔离、密钥脱敏、结构化响应校验。"],
        ["阶段二：内置模型扩展", "基于实测结果，逐步扩大自研 Qwen 在领域问答、实体关系抽取、分类与结构化生成任务中的使用。", "比较任务完成率、输出一致性、时延、资源占用与人工修改量。", "同一提示模板、同一 Schema、同一证据与审批链路。"],
        ["阶段三：统一协同", "在统一治理下，按数据敏感度、任务类型、时延与成本选择内置模型或云端服务。", "提升数据可控性和系统弹性，避免单一模型依赖。", "禁止绕过产品资格、客户授权、频控和发布审批。"],
    ]
    add_table(doc, model_rows, [1450, 3200, 2500, 2210])
    add_text(doc, "需要强调的是，双模型演进并不改变业务责任边界。无论模型来自本地还是云端，模型输出都只是包含事实、推断、建议和待确认事项的候选结果；活动发布、预算审批、正式本体更新和渠道触达必须由业务人员确认。", before=8, after=8)

    add_heading(doc, "3  实施方案", 1)
    add_heading(doc, "3.1  环境与可行性分析", 2)
    add_text(doc, "在工程实现方面，平台采用浏览器工作台、FastAPI 服务、PostgreSQL 数据库和 Docker Compose 的分层架构。浏览器负责展示机会、客群、产品、内容、审批、执行与复盘页面；服务端提供认证、租户隔离、数据接入、本体存储、模型配置和智能体接口；数据库承载结构化业务记录、知识片段、本体对象与关系。该架构降低了模块耦合度，便于后续接入经过授权的企业系统。")
    add_text(doc, "在数据可行性方面，系统支持从文件、接口和数据库接入业务资料，并通过对象类型、稳定标识、来源证据、置信度阈值和关系端点校验控制本体准入。数据源接入应使用授权范围内的只读账号、字段白名单、增量游标与脱敏规则，智能体不允许任意执行写入 SQL。")
    add_text(doc, "在 AI 可行性方面，当前模型服务采用 OpenAI 兼容协议接入阿里云百炼，并通过统一 Harness 调用知识检索、本体查询和活动上下文工具。项目自动化测试共 17 项通过，覆盖网络和连接恢复、输入边界、租户隔离、密钥保护、数据流水线、本体更新和六智能域治理；一次真实低流量端到端测试返回 HTTP 200，完成上下文加载、工具调用和综合回答，且保留人工审核门禁。")
    add_heading(doc, "3.1.1  推理服务与可靠性", 3)
    add_text(doc, "推理服务层通过统一模型配置管理百炼模型和内置自研模型。客户端兼容普通响应与流式响应：对流式场景逐段回传 token，并记录提示词、生成词和总 token 等用量字段；对非流式场景校验 choices/message/content 的响应结构，避免第三方服务返回异常格式时产生未处理错误。", before=0, after=6)
    add_text(doc, "仓库实现了最多 3 次的指数退避重试机制。针对 429、502、503、504 以及连接、读取、超时等瞬时异常，客户端会按策略重试；对于 400 等非瞬时错误则立即失败，防止无效重试消耗资源。系统对空问题和超过 12000 字符的输入在模型调用前拒绝，并在异常和 API 响应中避免暴露 API Key、Authorization 或 Bearer 认证信息。", before=0, after=6)
    add_text(doc, "这类工程指标体现的是推理系统的可用性与可治理性，而非营销收入本身。项目将模型可靠性、工具调用正确性、业务约束通过率与下游营销效果分层评估，避免用单一大模型分数替代完整的业务验证。", before=0, after=8)
    add_text(doc, "系统尚未把真实生产渠道回调、并发容量压测和所有外部数据连接器作为已完成成果。后续部署时，需在业务、数据安全和合规审批通过后接入生产数据源和发布活动。")
    add_heading(doc, "3.2  实施细节", 2)
    add_text(doc, "整体技术路线如下：")
    flow_rows = [
        ["阶段", "处理内容", "关键控制"],
        ["数据接入", "接收文件、授权 API 或数据库数据，登记来源、时间、租户与版本。", "格式校验、来源登记、权限范围与数据脱敏。"],
        ["知识与本体构建", "解析、清洗、切分、实体抽取、关系判断，形成知识片段和本体候选。", "证据溯源、置信度、对象/关系类型校验、人工准入。"],
        ["智能分析", "从机会到客群、产品、活动、内容和效果形成建议及解释。", "统一 Harness、工具轨迹、来源记录、模型输出结构校验。"],
        ["活动治理", "配置预算、客群、产品、渠道、内容、时间窗和执行批次。", "客户授权、频控、库存、资格、事实、敏感词与审批门禁。"],
        ["回流复盘", "关联送达、点击、出票、辅营购买、履约与投诉等聚合结果。", "归因结果可追溯，策略更新建议需人工确认。"],
    ]
    add_table(doc, flow_rows, [1700, 4300, 3360])
    add_text(doc, "平台的营销本体以“经营信号 - 营销机会 - 客户需求 - 客群快照 - 价值主张 - 产品包 - 策略方案 - 触点计划 - 内容与活动 - 审批 - 执行 - 反馈 - 归因 - 复盘学习”为核心链路。相较于单纯标签管理，该表达方式使营销建议能说明其业务对象、证据来源与约束条件，便于运营人员审核和复盘。")
    add_heading(doc, "3.2.1  本体论驱动的知识建模", 3)
    add_text(doc, "项目采用航司营销本体 v1.1，将业务语义编码为稳定对象类型、关系类型、可执行动作、业务函数和智能域契约。对象层覆盖 MarketSignal（市场信号）、Route（航线）、Flight（航班）、MetricObservation（经营指标观测）、CustomerNeed（客户需求）、AudienceSnapshot（客群快照）、ProductPackage（产品包）、Campaign（营销活动）、ExecutionBatch（执行批次）、Feedback（营销反馈）、AttributionResult（归因结果）、Review（效果复盘）、Evidence（业务证据）和 HumanDecision（人工决策）等核心对象。")
    add_text(doc, "关系层用于回答“为什么推荐、推荐给谁、依据什么执行、结果如何归因”四类问题。例如，市场信号通过 derived_from 关联原始证据，机会通过 concerns_route 关联航线、targets_audience 关联客群快照、uses_product_package 关联产品包；活动版本通过 requires_approval 关联审批任务、executes 关联执行批次、produces_feedback 关联渠道反馈；归因结果再通过 attributes_to 关联客群、产品、内容和渠道。这样，AI 的结论不再是脱离业务对象的自然语言，而是可以回到关系图谱逐项追溯。", before=0, after=6)
    add_text(doc, "本体还将‘事实、推断、建议、人工决定’分开存储。来自文档、接口或数据库的事实以 KnowledgeDocument、KnowledgeChunk、KnowledgeClaim 和 Evidence 保留来源定位；模型的结果以 Opportunity、Recommendation 等候选对象进入待确认状态；营销人员的采纳、修改或驳回以 HumanDecision 记录。该设计既为推理提供了可检索上下文，也避免将未经审核的模型输出直接写入正式经营知识。", before=0, after=8)
    add_heading(doc, "3.2.2  六智能域与业务函数", 3)
    domain_rows = [
        ["智能域", "主要推理输入", "关键输出与本体约束"],
        ["机会洞察", "市场信号、航线/航班、经营指标、历史复盘和规则。", "计算机会评分、识别经营异常，输出机会、营销目标、客户需求和证据。"],
        ["客群洞察", "聚合画像、需求、反馈、指标、客群属性和人工决策。", "生成可计算的客群快照，评估规模、吸引力、长期价值与频控。"],
        ["产品匹配", "机会、需求、客群快照、产品/产品包、航班和业务规则。", "校验库存、资格、衔接和适配性，形成待确认的产品匹配建议。"],
        ["活动编排", "营销目标、客群、产品包、价值主张、渠道、预算和规则。", "形成策略与触点计划；需要审批的动作不能由模型直接发布。"],
        ["内容生成", "已审核产品事实、客群偏好、渠道规范和敏感词规则。", "生成内容资产并通过事实、品牌、合规和渠道格式检查。"],
        ["效果分析", "执行批次、触达、点击、交易、履约、投诉和指标观测。", "计算活动效果、增量归因和复盘建议，规则更新仍需人工确认。"],
    ]
    add_table(doc, domain_rows, [1550, 3550, 4260])
    add_text(doc, "每个智能域均通过统一 Harness 运行。一次运行至少包含业务与本体上下文加载、模型提供者选择、治理检查、人工审核判断和运行结束事件。代码层的智能域契约限定了各域能够读取哪些对象、可以生成哪些候选对象、能够调用哪些业务函数，从机制上减少模型越权读取或越过业务节点直接执行的风险。", before=8, after=8)
    img = ROOT / "apps" / "web-v32" / "v32-final.png"
    if img.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(img), width=Cm(15.5))
        add_caption(doc, "图 1  东方航空智能营销平台原型工作台（项目现有界面截图）")
    add_text(doc, "原型工作台围绕机会、客群、产品、内容、审批、执行和效果复盘组织业务操作，并提供营销助手入口。营销助手先读取允许访问的营销知识、本体关系和活动上下文，再基于工具观测给出可执行建议、数据依据、风险项和需要人工确认的事项。")
    ont = ROOT / "docs" / "screenshots" / "ontology-workbench-v1.0.png"
    if ont.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(ont), width=Cm(13.0))
        add_caption(doc, "图 2  营销知识与本体工作台（项目现有界面截图）")
    add_heading(doc, "3.3  AI 推理与营销本体协同", 2)
    add_text(doc, "在用户侧，营销助手不是直接对数据库进行任意查询的聊天机器人。它先接收营销问题，加载当前租户内被授权的知识、本体、活动和工具上下文，再由模型规划下一步业务工具。工具返回的机会、客群、产品、活动或反馈数据被重新纳入上下文，模型据此形成含业务依据、风险和待人工确认事项的综合答复。运行轨迹向用户展示上下文加载、规划、模型调用、工具调用和最终生成等步骤。")
    add_text(doc, "在数据侧，文件导入遵循“接收 - 解析 - AI 抽取 - 校验 - 入库”的处理链。系统先对文件格式、租户和来源进行检查，再抽取知识片段、识别业务对象和关系、进行语义校验。候选本体更新必须由业务人员确认后才会写入正式图谱；证据不足的信息可以仅保留在知识层，避免把低置信度内容错误传播到下游活动决策。")
    inference_rows = [
        ["推理环节", "AI 技术作用", "本体与治理约束", "输出效果"],
        ["问题理解", "识别机会、客群、产品、活动或复盘意图，选择对应智能域。", "限定在当前租户和授权数据范围内。", "把自然语言问题映射为业务对象与可调用工具。"],
        ["上下文构建", "检索知识片段、本体对象、关系和活动状态，形成可追溯上下文。", "对象携带来源、证据、置信度、有效期和版本。", "减少无依据生成，支持回答中说明数据来源。"],
        ["工具编排", "模型规划并调用机会评分、客群评估、资格校验、频控、效果计算等函数。", "智能域契约限制可读对象、可写候选和可调用函数。", "将通用语言模型能力落到可验证的业务计算。"],
        ["结构化生成", "生成机会解释、产品建议、渠道内容、审批摘要或复盘建议。", "输出经过 Schema、事实、敏感词和业务规则校验。", "形成可编辑、可审核、可版本化的业务资产。"],
        ["人工闭环", "根据反馈归纳下一轮客群、内容、时机和渠道建议。", "发布、审批、本体更新和规则生效均保留人工确认。", "使 AI 推理结果持续受业务结果校验。"],
    ]
    add_table(doc, inference_rows, [1400, 2700, 2900, 2360])
    add_text(doc, "模型效果的衡量被分为三层：第一层是推理工程指标，包括响应结构合法性、重试恢复、token 用量、耗时、工具调用次数和失败率；第二层是业务中间指标，包括机会置信度、客群规模、产品适配率、内容事实校验和审批通过情况；第三层是营销结果指标，包括送达、回执、点击、出票、辅营购买、增量收入和 ROI。三层指标共同用于判断模型是否真正提升了营销工作质量。", before=8, after=8)
    add_heading(doc, "3.4  计划与分工", 2)
    add_text(doc, "项目以原型完善、数据接入验证、智能体治理、场景演示和材料沉淀为主线推进。以下为当前竞赛阶段的建议分工，团队可根据实际开发进展进行调整。")
    plan_rows = [
        ["阶段", "主要任务", "预期产出"],
        ["需求与建模", "梳理航空营销闭环、业务对象、数据边界和人工门禁。", "业务本体、数据字段与流程说明。"],
        ["系统开发", "完善工作台、服务接口、数据库模型、数据处理与智能体链路。", "可运行原型、接口与部署脚本。"],
        ["验证与治理", "完成异常场景测试、权限/隔离检查、模型接入和低流量端到端验证。", "测试记录、风险清单与治理边界。"],
        ["竞赛展示", "制作典型营销场景演示，整理项目文档、视频与可复现工程材料。", "参赛项目文档、演示视频与工程压缩包。"],
    ]
    add_table(doc, plan_rows, [1800, 4400, 3160])
    division_rows = [
        ["成员", "建议职责"],
        ["胡于飞", "项目统筹、系统架构、后端接口与 AI 智能体治理。"],
        ["林李洋", "前端工作台、业务交互、活动流程与演示场景设计。"],
        ["张婷婷", "数据资料整理、业务验证、测试记录与竞赛材料编制。"],
    ]
    add_table(doc, division_rows, [2200, 7160])

    add_heading(doc, "3.5  验证指标与典型场景", 2)
    add_text(doc, "项目以“上海—三亚国庆早鸟”作为原型营销场景：机会洞察将目的地旅游热度、App 搜索增长、航线供给和预售窗口关联为机会候选；客群洞察形成‘近期搜索、处于比价期且未出票’的聚合客群快照；产品匹配将早鸟机票、行李、优选座位和目的地权益组合为产品包；内容生成按 App、短信等渠道输出版本；审批与执行模块再校验预算、授权、频控、库存、资格和内容合规。")
    add_text(doc, "在当前已登录的原型工作台中，该场景展示的 AI 和业务中间指标包括：机会清单共 18 条，‘上海—三亚国庆早鸟’机会价值评分为 92，本体对象置信度为 94%；目标客群快照为 36,420 人，已完成关系扩展和营销疲劳排除；产品包适配率为 91.6%。这些数值用于展示原型的推理、匹配和本体关联效果，不应解释为已在生产环境中实现的经营承诺。")
    metrics_rows = [
        ["指标类别", "原型已观测指标", "AI 与本体论的作用", "使用边界"],
        ["机会推理", "18 条机会；上海—三亚机会评分 92；对象置信度 94%。", "融合市场、航线、客群和证据对象，形成可追溯机会候选。", "评分用于排序与解释，正式机会仍需人工确认。"],
        ["客群与产品", "目标客群 36,420 人；产品适配率 91.6%。", "根据旅程阶段、行为、产品资格和保护规则完成匹配。", "仅使用聚合画像，避免暴露个人旅客明细。"],
        ["内容与审批", "App/短信/微信等内容版本；审批按时完成率 96.8%。", "生成渠道内容，并用产品事实、频控、敏感词和预算规则约束。", "内容及活动发布必须通过人工审核。"],
        ["执行与回流", "成功触达率 98.6%；回执回流率 96.8%；失败待补偿 126 人。", "按执行批次、渠道和反馈建立可归因对象关系。", "为原型工作台运行数据，需在真实渠道接入后持续验证。"],
        ["营销复盘", "点击率 11.4%（较基线 +2.3pp）；出票转化率 9.8%（+1.8pp）；ROI 4.6。", "把内容、客群、产品、渠道与结果关联，生成下一轮优化建议。", "属于演示场景看板指标，不作为生产实际收益披露。"],
        ["推理可靠性", "自动化稳健性测试报告记录 17 项通过；支持最多 3 次瞬时错误重试。", "验证模型接入、结构化响应、租户隔离、密钥保护和六智能域门禁。", "当前容器运行镜像未包含 pytest 命令，现场未重复执行全量测试。"],
    ]
    add_table(doc, metrics_rows, [1450, 2700, 3000, 2210])
    add_text(doc, "典型场景的价值不在于单次生成营销文案，而在于模型、规则和本体共同驱动的闭环：模型从证据和关系中发现机会，业务函数完成规模、资格、频控和效果计算，运营人员审核并执行，执行结果再回写为反馈、归因和复盘对象。这样可将每轮模型推理同实际营销结果关联，为内置 Qwen 模型后续扩大使用提供统一的对照测试和持续评估基础。", before=8, after=8)

    add_heading(doc, "3.6  平台功能与业务工作流", 2)
    add_text(doc, "平台已将航空营销从数据接入到效果复盘的主要工作节点组织为可操作的工作台模块。每个模块都保留业务人员可理解的名称、状态和证据，而将外部系统标识、本体类型编码和关系端点封装在服务层。AI 并非独立悬置于业务之外，而是在各模块之间传递对象、证据、规则与待确认事项。", before=0, after=7)
    module_rows = [
        ["功能模块", "面向用户的操作", "AI/本体能力", "输出与人工边界"],
        ["机会洞察", "查看市场热点、航线表现和机会清单，进入某一机会查看证据。", "关联市场信号、航线、航班、经营指标和历史复盘，生成机会解释与评分。", "输出机会候选、影响范围和建议时机；是否进入营销策划由人工确认。"],
        ["客群洞察", "选择画像条件、查看客群规模、旅程特征和保护排除。", "将需求、画像、历史反馈、授权和频控规则连接为版本化客群快照。", "输出聚合人群与解释；不在营销界面展示不必要的个人明细。"],
        ["产品与权益", "检索客票、辅营、卡券、会员权益和产品组合。", "依据航班、库存、资格、有效期、运价与权益规则进行适配和冲突提示。", "输出主推与替代产品建议；最终产品承诺须由业务人员审核。"],
        ["内容工作台", "编辑短信、App、微信、官网等渠道内容版本。", "基于已审核产品事实、价值主张、客群语境和渠道规范生成候选内容。", "输出可编辑内容与风险提示；发布前执行事实、品牌、敏感词和合规检查。"],
        ["活动与审批", "配置目标、预算、客群、产品、渠道、频控、时间窗并提交审批。", "把活动对象与策略、触点计划、审批、执行批次和规则关联。", "未通过预算、授权、库存、资格和审批校验的活动不能进入发布状态。"],
        ["执行监测", "查看投放批次、送达、失败、回执和异常待补偿。", "将渠道状态写入执行批次和反馈对象，形成可归因事件链。", "输出执行状态；真实渠道接入与补偿策略需在生产环境持续验证。"],
        ["效果复盘", "查看触达、点击、购票、核销、辅营、收入和投诉等聚合指标。", "将结果归因到客群、产品包、内容、渠道和策略，生成异常解释与优化建议。", "输出复盘与下一轮建议；规则更新仍需人工确认。"],
        ["知识与本体", "查看文件处理进度、候选对象关系、证据和图谱工作台。", "完成知识切分、实体关系抽取、语义校验、置信度计算和版本留痕。", "候选更新经确认后才写入正式图谱，低置信度信息可只保留知识层。"],
        ["模型与审计", "配置提供商、查看可用模型、连接状态、调用次数与 Token 用量。", "统一管理百炼与内置测试模型，记录模型、工具、上下文和治理事件。", "密钥只在服务端保存；管理员可管理配置，普通业务角色不能越权访问。"],
    ]
    add_table(doc, module_rows, [1500, 2650, 2700, 2510])
    add_text(doc, "上述模块在原型工作台中形成明确的使用顺序：先通过机会洞察识别经营窗口，再通过客群与产品模块形成可执行的供给匹配，随后在内容、活动、审批与执行模块中完成受控触达，最后由效果复盘把结果回写为下一轮 AI 推理的证据和业务上下文。", before=8, after=8)

    add_heading(doc, "3.7  数据处理与知识本体构建", 2)
    add_text(doc, "航空营销信息存在于航线调整通知、服务规则、运价资料、经营报表、产品权益、渠道日志和复盘材料等不同形态的数据中。平台将文件、授权 API 和数据库记录纳入统一处理链，目标不是把原始内容简单转成文本，而是让每一项可被营销决策使用的事实都能够回答“来自哪里、何时生效、适用于什么对象、可信程度如何”。")
    add_heading(doc, "3.7.1  数据接入、解析与候选审核", 3)
    pipeline_rows = [
        ["处理阶段", "系统处理内容", "AI 推理作用", "质量与治理控制"],
        ["接收检查", "创建任务并记录文件名、格式、来源、租户、提交人和时间。", "按文档、规则、产品、经营指标、渠道结果等类别识别处理策略。", "检查格式、大小、租户归属、重复文件和授权范围。"],
        ["结构解析", "抽取正文、章节、页码、表格、图片与来源位置；复杂文件可调用解析服务。", "识别标题、表头、字段、段落和语义区块，为后续抽取保留结构。", "保留原始文件哈希和解析版本；低质量 OCR 或错位表格进入审核队列。"],
        ["知识切分", "按标题、页码、段落和表格行列切分知识片段。", "生成适合检索和引用的片段摘要及业务关键词。", "每个片段保留来源定位，避免答案无法回溯到原始材料。"],
        ["对象关系抽取", "形成航线、航班、产品、规则、机会、客群、活动和渠道等候选对象。", "从文本和表格中识别属性、适用条件、时间窗与对象关系。", "按对象类型和关系端点校验，区分事实、推断与建议。"],
        ["语义校验", "检查航班条件、库存、资格、有效期、MCT、渠道和规则冲突。", "对不完整、冲突或低置信度信息生成待确认说明。", "高影响事实不自动覆盖；按来源优先级、有效期与人工确认决定准入。"],
        ["知识与本体更新", "保存知识文档、知识片段、候选对象、关系、证据和审核状态。", "通知六智能域可用的上下文范围，支持后续检索和推理。", "只有确认后的候选才更新正式图谱；保留版本、操作者与撤销依据。"],
    ]
    add_table(doc, pipeline_rows, [1450, 2600, 2700, 2610])
    add_text(doc, "工作台将处理链显示为可追踪状态，例如排队、接收检查、解析结构、解析完成、知识入库、智能体处理、语义校验、待人工确认和本体已更新。业务人员可以在任务详情中查看候选对象数、候选关系数、模型任务、运行事件和输出预览。这一设计使 AI 抽取不再是黑箱批处理，而是可检查、可驳回、可重新执行的业务流程。", before=8, after=8)
    add_heading(doc, "3.7.2  营销本体的对象、关系与状态", 3)
    ontology_rows = [
        ["语义层", "典型对象或关系", "对 AI 推理的作用"],
        ["经营与机会", "MarketSignal、Route、Flight、MetricObservation、Opportunity。", "将热点、供给、客座率、预售窗口与历史结果组织为可解释的机会证据链。"],
        ["需求与客群", "CustomerNeed、CustomerAggregate、AudienceSnapshot、HumanDecision。", "把聚合画像、行为、需求、授权、保护规则和人工确认固化为可复现客群快照。"],
        ["产品与价值", "Product、ProductPackage、ValueProposition、BusinessRule。", "让模型在事实、资格、库存、有效期和权益边界内匹配产品组合。"],
        ["策略与执行", "StrategyPlan、TouchpointPlan、Campaign、ApprovalTask、ExecutionBatch。", "把自然语言建议转成带渠道、节奏、预算、频控和审批的可治理计划。"],
        ["反馈与学习", "Feedback、AttributionResult、Review、Recommendation。", "将送达、互动、交易、履约和投诉关联到策略要素，形成下一轮改进依据。"],
        ["证据与溯源", "KnowledgeDocument、KnowledgeChunk、KnowledgeClaim、Evidence。", "支持模型答案引用来源片段，区分原始事实、候选推断和人工决策。"],
    ]
    add_table(doc, ontology_rows, [1800, 3500, 4060])
    add_text(doc, "本体关系并不是对数据表字段的简单重命名。比如 Opportunity 通过 concerns_route 关联航线、targets_audience 关联客群快照、uses_product_package 关联产品包；Campaign 通过 requires_approval 关联审批任务、executes 关联执行批次、produces_feedback 关联反馈；AttributionResult 再通过 attributes_to 回到客群、内容、产品与渠道。模型每次生成建议时可以沿这些关系检索所需上下文，并向业务人员说明推荐依据。", before=8, after=8)

    add_heading(doc, "3.8  推理编排、模型路由与可观测性", 2)
    add_text(doc, "平台的 AI 运行采用统一 Harness，而不是由每个页面分别调用大模型。Harness 将营销问题、当前租户、角色权限、知识检索、本体关系、活动状态、工具函数和模型配置组织为一次受控运行。在输出最终建议之前，系统记录上下文加载、提供者选择、治理检查、模型调用、工具轨迹、人工审核判断与结束状态。")
    add_heading(doc, "3.8.1  一次营销推理的运行链", 3)
    inference_steps = [
        ["步骤", "运行内容", "可观测证据"],
        ["1. 请求校验", "校验登录身份、租户、角色、输入长度和可用模型提供者。", "请求标识、租户标识、拒绝原因和安全审计。"],
        ["2. 上下文加载", "读取授权范围内的知识片段、本体对象关系、活动状态与业务规则。", "知识来源、图谱对象、版本、有效期和上下文加载事件。"],
        ["3. 智能域选择", "将问题映射到机会洞察、客群洞察、产品匹配、活动编排、内容生成或效果分析。", "智能域契约、允许读取的对象和允许调用的函数。"],
        ["4. 模型与工具编排", "由模型生成受约束的下一步分析，并调用机会评分、资格校验、频控、效果计算等工具。", "模型名称、提示词/生成 Token、工具名称、输入输出摘要和耗时。"],
        ["5. 结构化校验", "检查模型响应结构、事实引用、敏感词、权限边界和业务规则。", "Schema 校验结果、风险项、失败或降级事件。"],
        ["6. 人工门禁", "对活动发布、内容对外发送、本体正式更新、预算和规则生效判断是否需要确认。", "待确认事项、审批状态、人工采纳/修改/驳回记录。"],
        ["7. 结果回写", "保存运行记录、候选结果、证据、反馈和复盘关联。", "AgentRun、Recommendation、HumanDecision 与 Review 对象关系。"],
    ]
    add_table(doc, inference_steps, [1000, 4550, 3810])
    add_heading(doc, "3.8.2  模型路由与双模型使用策略", 3)
    add_text(doc, "前期暂用阿里云百炼 AgentStudio 的 qwen3.7-flash 作为默认推理服务，优先验证营销问答、知识检索、本体查询、工具调用、流式输出和人工审核门禁在真实原型中的连通性。团队内置测试模型对应本地自研 Qwen 模型的受控验证入口，用于在同一 Harness、同一提示模板、同一输出 Schema 和同一业务门禁下进行流程对照。")
    routing_rows = [
        ["路由维度", "百炼主用阶段", "内置 Qwen 扩展阶段", "统一约束"],
        ["任务类型", "复杂业务问答、跨对象归纳、活动建议和通用文本生成。", "领域问答、对象关系抽取、分类、结构化生成和可本地执行的辅助任务。", "使用同一智能域契约、工具白名单和输出结构。"],
        ["数据敏感度", "使用已授权、脱敏和聚合后的上下文。", "对需要更高本地可控性的合规场景逐步扩大验证范围。", "不得绕过数据最小化、租户隔离和权限控制。"],
        ["质量与性能", "以任务完成、证据绑定、人工修改量、时延和稳定性验证端到端链路。", "在相同测试集上与云端服务对照，逐步扩大通过门槛的任务比例。", "每次调用记录模型、Token、耗时、重试、失败和人工反馈。"],
        ["业务责任", "模型输出为候选解释、建议和待确认事项。", "模型输出同样受事实、规则、审批和发布门禁约束。", "活动发布、本体更新、预算与渠道触达由人负责。"],
    ]
    add_table(doc, routing_rows, [1500, 2700, 2700, 2460])
    add_text(doc, "后续扩大内置模型使用并非简单替换模型名称，而是以任务级对照为依据：在同一知识和本体上下文、同一问题集、同一工具约束和同一人工审核规则下，比较结构化输出合法率、证据引用完整率、业务规则通过率、人工修改量、端到端时延和资源占用。只有满足预设质量与安全门槛的任务，才逐步由内置模型承担更多比例。", before=8, after=8)

    add_heading(doc, "3.9  AI 评测体系与稳健性验证", 2)
    add_text(doc, "本项目不以单一大模型主观评分作为效果结论，而是采用‘推理工程质量 - 业务中间质量 - 营销结果质量’三级指标体系。前两层用于判断 AI 能否在约束内可靠工作，第三层用于判断营销方案在真实授权渠道接入后是否带来可验证的业务改善。原型看板数值只用于展示当前系统状态和演示场景，不替代生产经营统计。")
    eval_rows = [
        ["评测层级", "核心指标", "评测对象", "判定目的"],
        ["推理工程质量", "响应结构合法率、重试恢复、错误率、Token 用量、时延、工具调用轨迹。", "模型客户端、Harness、工具编排和运行事件。", "确认模型服务可用、可观测、出现异常时可控。"],
        ["知识与本体质量", "对象抽取准确性、关系端点合法性、证据覆盖、置信度、冲突和过期状态。", "数据处理任务、候选对象关系、知识片段与审核记录。", "确认推理上下文可追溯，避免低质量事实进入正式图谱。"],
        ["业务建议质量", "机会解释完整性、客群规模合理性、产品资格通过率、内容事实校验、审批通过情况。", "六智能域的候选输出、人工修改/采纳/驳回记录。", "确认 AI 建议与航空营销规则一致、可被业务人员使用。"],
        ["营销结果质量", "送达、回执、点击、购票、核销、辅营购买、投诉、增量效果与 ROI。", "执行批次、渠道反馈、交易履约和归因结果。", "确认活动效果而非仅确认文本生成效果。"],
    ]
    add_table(doc, eval_rows, [1700, 3000, 2700, 1960])
    test_rows = [
        ["测试场景", "仓库实现或验证点", "对 AI Inference 的意义"],
        ["瞬时异常恢复", "对 429、502、503、504、连接、读取和超时异常实施最多 3 次指数退避重试。", "避免短暂网络波动造成一次推理链直接失败，并保留可诊断错误。"],
        ["响应结构校验", "校验非流式 choices/message/content；流式输出逐段收集 token 和用量。", "防止外部模型返回异常格式时将错误结果进入业务流程。"],
        ["输入与密钥保护", "拒绝空问题和超过 12000 字符输入；异常和 API 响应不暴露 API Key、Authorization、Bearer。", "降低提示注入、无效消耗和凭证泄露风险。"],
        ["租户与提供者隔离", "验证未知提供者、跨租户提供者访问被拒绝，并检查输出不泄露提供商密钥。", "确保模型调用、知识上下文和审计记录遵守租户边界。"],
        ["六智能域门禁", "验证六个智能域都经历治理检查、上下文加载、提供者选择、人工审核和运行结束事件。", "保证模型输出不能绕过活动审批、本体更新和发布控制。"],
    ]
    add_table(doc, test_rows, [1700, 4200, 3460])
    add_text(doc, "项目现有自动化稳健性测试报告记录 17 项通过，覆盖模型重试、异常响应、输入边界、密钥保护、租户隔离、数据管道、本体更新和六智能域治理。该结论说明原型已对若干关键推理风险建立工程防护；但不等同于已完成真实生产渠道的全量压力测试、经营效果检验或所有外部连接器的生产验收。", before=8, after=8)

    add_heading(doc, "3.10  治理边界与后续建设路线", 2)
    add_text(doc, "航空营销涉及旅客权益、价格、库存、会员资格、触达同意、频控、预算和品牌合规，AI 的价值必须建立在责任边界清晰的前提上。平台把‘模型可以生成什么’与‘系统可以执行什么’分离：模型生成事实摘要、候选关系、建议与风险说明；系统执行的任何对外动作都必须经过规则服务、权限校验和人工审批。")
    governance_rows = [
        ["治理维度", "当前原型控制", "后续生产化建设重点"],
        ["数据与隐私", "使用租户隔离、聚合画像、来源记录、字段白名单和服务端密钥管理。", "接入企业数据分级分类、脱敏策略、访问审计、数据保留与删除机制。"],
        ["事实与知识", "候选对象关系保留证据、置信度、状态和人工确认；支持知识层与正式图谱分层。", "完善来源优先级、冲突仲裁、有效期失效、版本回滚和批量审核工作流。"],
        ["模型与推理", "统一模型提供商配置、Token 统计、重试、响应校验和 Harness 运行轨迹。", "建立模型基准集、任务路由阈值、成本监测、漂移监控和内置模型灰度策略。"],
        ["营销动作", "审批、预算、资格、库存、频控、敏感词和内容事实作为活动门禁。", "与正式渠道、OA 审批、预算系统和异常补偿系统完成授权对接。"],
        ["运营与复盘", "展示执行、回执、转化和复盘对象关系，原型指标用于演示。", "建立标准归因口径、对照实验、增量评估、持续复盘和策略版本治理。"],
    ]
    add_table(doc, governance_rows, [1600, 3700, 4060])
    add_text(doc, "近期，项目将继续补齐数据处理任务的可视化溯源、候选审核、智能域演示和典型活动闭环。验收时应确保每一项建议均可查看来源、本体关系、运行轨迹和人工确认状态。", before=8, after=5)
    add_text(doc, "中期，团队将在统一评测集和同一业务门禁下，将通过质量、安全、时延和人工修改量阈值的领域抽取、分类与结构化任务，逐步切换或路由至内置 Qwen 自研模型；任一任务均应支持回退至云端服务。", before=0, after=5)
    add_text(doc, "后期，平台将在授权条件下接入企业数据源、审批和渠道回调，形成可审计运营闭环；只有完成数据安全、业务、合规和运维验收后，才开放正式发布能力。", before=0, after=8)
    add_text(doc, "综上，项目的核心竞争力不在于单独调用某一个大模型，而在于把 AI Inference、营销本体、业务函数、数据证据、人工决策和营销结果编排为同一条可追溯链路。前期以百炼服务支撑端到端验证，后续以可量化对照为基础逐步扩大内置 Qwen 自研模型使用，最终形成兼顾推理质量、数据可控性、业务合规与持续学习能力的航空智能营销平台。", before=8, after=8)

    add_heading(doc, "3.10.1  典型场景 AI 推理执行清单", 3)
    add_text(doc, "以‘上海—三亚国庆早鸟’原型场景为例，系统并非在得到一句自然语言指令后直接生成文案，而是按以下受控链条推进。各步骤均保留可检查输入、候选输出和人工决策点，便于将 AI Inference 的作用与实际营销运行环节对应。", before=0, after=6)
    scenario_rows = [
        ["环节", "AI 与本体处理", "可观测输出", "人工确认点"],
        ["机会识别", "关联目的地热度、App 搜索增长、航线供给、客座率观测和预售窗口，生成 Opportunity。", "机会评分、证据来源、影响航线、时间窗与对象置信度。", "确认是否具备营销价值并进入活动策划。"],
        ["客群快照", "将近期搜索、比价期、未出票、旅程特征、授权与营销疲劳排除组合为 AudienceSnapshot。", "聚合规模、标签解释、排除规则、有效期和可触达范围。", "确认客群口径、保护名单与频控条件。"],
        ["产品匹配", "依据航班、产品包、行李、优选座位、目的地权益、库存、资格与有效期进行规则核验。", "主推产品、替代方案、适配率、不可用原因与事实证据。", "确认对外可承诺的产品组合和权益文案。"],
        ["内容与活动", "按 App、短信等渠道规格生成内容候选，并将目标、预算、渠道、时间窗和频控组织为 Campaign。", "内容版本、事实校验、敏感词提示、触点计划与审批摘要。", "审核内容、预算、渠道和发布节奏。"],
        ["执行与回流", "把授权渠道的送达、回执、点击和失败状态关联为 ExecutionBatch 与 Feedback。", "执行批次、送达率、回执率、失败待补偿和渠道状态。", "确认异常补偿与真实渠道数据的归因口径。"],
        ["复盘学习", "将内容、客群、产品包、渠道、交易履约与投诉结果连接为 AttributionResult 和 Review。", "漏斗指标、异常解释、人工采纳记录与下一轮 Recommendation。", "确认策略、规则或内容模板是否更新。"],
    ]
    add_table(doc, scenario_rows, [1350, 3200, 2700, 2110])

    doc.add_page_break()
    add_heading(doc, "4  参考资料", 1)
    references = [
        "[1] 《东方航空智能营销平台》项目 README，项目仓库，v2.7。",
        "[2] 《东方航空智能营销平台数据处理与 AI 营销全流程说明 v1.0》，项目技术文档。",
        "[3] 《航司营销业务本体 v1.1》，项目技术文档。",
        "[4] 《市场热点智能体与本体处理方案 v1.0》，项目技术文档。",
        "[5] 《AI 智能体稳健性测试报告 v1.0》，项目测试文档。",
        "[6] 第八届中国研究生人工智能创新大赛初赛作品提交规范。",
        "[7] 第八届中国研究生人工智能创新大赛项目文档模板。",
        "[8] GitHub：PLiO-LIN/CeairMarketing，项目主分支代码与提交记录。",
        "[9] services/platform-api/app/ontology/semantic_model.py，航司营销本体 v1.1 的对象、关系、动作、函数与智能域契约实现。",
        "[10] services/platform-api/app/llm.py 与 tests/test_llm_robustness.py，模型调用、流式响应、重试和安全测试实现。",
    ]
    for ref in references:
        add_text(doc, ref, size=10.5, before=0, after=6,
                 align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], first_page=True)
    core = doc.core_properties
    core.title = "东方航空智能营销平台项目文档"
    core.subject = "第八届中国研究生人工智能创新大赛"
    core.author = "丽娃谷风"
    core.comments = "项目文档 V1.3"
    cover(doc)
    toc(doc)
    revision_history(doc)
    main_body(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
